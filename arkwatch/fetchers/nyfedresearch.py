"""nyfedresearch.py — NY Fed research & survey datasets (www.newyorkfed.org files).

8 datasets (endpoints live-verified 2026-09-19 by the 8-agent research
workflow; calibration anchors recorded in docs/sources.md):
  HHDC   Quarterly Report on Household Debt & Credit (databank XLSX, quarterly)
  SCE    Survey of Consumer Expectations, core module (XLSX, monthly)
  GSCPI  Global Supply Chain Pressure Index (interactive CSV, monthly)
  ESMS   Empire State Manufacturing Survey (SA diffusion CSV, monthly)
  ACM    Adrian-Crump-Moench term premia (legacy .xls, daily)
  HPW    Heise-Pearce-Weber labor-market tightness (chart CSV, monthly)
  LW     Laubach-Williams r-star + HLW variant (XLSX, quarterly)
  MCT    Multivariate Core Trend inflation (chart CSV, monthly)

Verified pitfalls:
  - HHDC 'Total' cells are FORMULAS — openpyxl data_only=True is mandatory or
    the stored "value" is the formula string. Balance sheet (Page 3) is in
    $TRILLIONS while originations sheets (Page 6/8) are $BN: ×1000 happens at
    ONE site (_hhdc). Quarter labels are 'YY:Qn' TEXT.
  - GSCPI's downloads-page file is legacy OLE2 .xls despite the .xlsx name —
    we use the interactive vintage-matrix CSV instead (plain text, also carries
    the 4 months of 1997 the xls lacks); CURRENT vintage = LAST column.
  - ESMS: in the allseries file the suffix NSA means 'No-change Share, SA' —
    NOT 'not seasonally adjusted'; only the diffusion CSV is ingested.
    surveyDate stamps the survey month's END (2026-09-30 for the survey
    released Sep 15) — normalized to month START (report_month convention).
  - ACM ACMTermPremium.xls is a ~10MB legacy .xls — xlrd only (openpyxl cannot
    read OLE2); all cells arrive as TEXT (dates '17-Sep-2026', values strings).
    Process-cached: one download serves all 4 series per run.
  - SCE sheets: labels in row index 3, data from row 4 — map columns BY LABEL
    (job-loss and quit-probability share one sheet; position guesses swap them).
  - HPW (and its ECI column) are z-SCORES normalized to mean 0 / sd 1 — NOT
    percent; the trailing month row is an empty placeholder — dropped.
  - LW/HLW: the whole 1961-present history is RE-ESTIMATED every quarterly
    release (never append-only-trust); one-sided estimates are the first block
    (rstar,g,z,gap at cols C-F), two-sided at H-K (terminal rows coincide —
    not duplicates); prose rows + 2 header rows precede data.
  - MCT CSV: 4 header rows, leading unnamed column, FOUR columns all named
    'MCT' = band-lo | point | band-hi | vs-2017-19-avg (semantics verified
    numerically across all 799 rows; enforced at parse: lo<=point<=hi).
    Sector columns are CONTRIBUTIONS in pp — not inflation rates.
  - Revisions: HHDC/MCT/LW/GSCPI/HPW rewrite history — realtime rows are
    updated with vintage snapshots by the harvest (same contract as FRED).

Routing: registry family stays NYFED: — nyfed.fetch_latest delegates here via
knows(); the 06:00 harvest calls fetch_window (gap-heal) and the Sunday
backfill re-ingests full history for revision pickup.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime

import requests

BASE_WWW = "https://www.newyorkfed.org"
URLS = {
    "sce": f"{BASE_WWW}/medialibrary/interactives/sce/sce/downloads/data/frbny-sce-data.xlsx",
    "gscpi": f"{BASE_WWW}/medialibrary/research/interactives/data/gscpi/gscpi_interactive_data.csv",
    "empire": f"{BASE_WWW}/medialibrary/media/survey/empire/data/esms_seasonallyadjusted_diffusion.csv",
    "acm": f"{BASE_WWW}/medialibrary/media/research/data_indicators/ACMTermPremium.xls",
    "hpw": f"{BASE_WWW}/medialibrary/Research/Interactives/Data/labor-market-tightness/HPW_monthly_share.csv",
    "lw": f"{BASE_WWW}/medialibrary/media/research/economists/williams/data/"
    "Laubach_Williams_current_estimates.xlsx",
    "hlw": f"{BASE_WWW}/medialibrary/media/research/economists/williams/data/"
    "Holston_Laubach_Williams_current_estimates.xlsx",
    "mct": f"{BASE_WWW}/medialibrary/Research/Interactives/Data/mct/mct-chart-data.csv",
    "hhdc_index": f"{BASE_WWW}/microeconomics/databank.html",
}


class NyFedResearchError(RuntimeError):
    pass


# --- transport -------------------------------------------------------------------

_bytes_cache: dict[str, bytes] = {}


def _download(key: str) -> bytes:
    """Fetch a URL (process-cached — the 10MB ACM file serves all 4 series)."""
    if key in _bytes_cache:
        return _bytes_cache[key]
    r = requests.get(URLS[key], timeout=(10, 120))
    if r.status_code != 200 or not r.content:
        raise NyFedResearchError(f"nyfedresearch {key}: HTTP {r.status_code} / empty")
    _bytes_cache[key] = r.content
    return r.content


# --- small date helpers (single conversion sites) ---------------------------------


def _dd_mon_yyyy(s: str) -> str:
    """'17-Sep-2026' / '31-Aug-2026' → ISO."""
    return datetime.strptime(s.strip(), "%d-%b-%Y").date().isoformat()


def _m_dy(s: str) -> str:
    """'7/1/2026' (M/D/YYYY, the chart-feed convention) → ISO."""
    return datetime.strptime(s.strip(), "%m/%d/%Y").date().isoformat()


def _month_start(month_end_iso: str) -> str:
    """'2026-09-30' (ESMS survey stamp) → '2026-09-01' (report_month)."""
    return month_end_iso[:8] + "01"


def _is_quarter_label(v) -> bool:
    """True for '26:Q2' text / datetime cells — False for footnote prose
    ('* 2026Q2 report includes a revision…' rides below the data, live 2026-09)."""
    return isinstance(v, datetime) or (
        isinstance(v, str) and re.match(r"^\s*\d{2,4}\s*[:\-]?\s*Q\s*\d\s*$", v, re.IGNORECASE)
    )


def _quarter_start(label: str) -> str:
    """HHDC '26:Q2' / datetime / ISO text → quarter-start ISO ('2026-04-01')."""
    if isinstance(label, datetime):
        return f"{label.year}-{(label.month - 1) // 3 * 3 + 1:02d}-01"
    s = str(label).strip()
    m = re.match(r"(\d{2,4})[:\s-]*Q(\d)", s, re.IGNORECASE)
    if m:
        yy = int(m.group(1))
        # century pivot: '26:Q2'→2026 (yy<50 → 2000s), '99:Q4'→1999
        # (live-caught inverted twice in a row: the probe kept storing
        # 1926-04-01 — the digit class this repo burns on)
        if yy < 50:
            year = 2000 + yy
        elif yy < 100:
            year = 1900 + yy
        else:
            year = yy
        return f"{year}-{(int(m.group(2)) - 1) * 3 + 1:02d}-01"
    m2 = re.match(r"(\d{4})-(\d{2})", s)  # already ISO-ish → snap to quarter start
    if m2:
        return f"{m2.group(1)}-{(int(m2.group(2)) - 1) // 3 * 3 + 1:02d}-01"
    raise NyFedResearchError(f"hhdc: unparseable quarter label {label!r}")


def _f(v) -> float | None:
    """Cell → float, tolerating text numbers and blanks (returns None)."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


# --- HHDC (quarterly household debt) ----------------------------------------------


def _hhdc() -> dict[str, list[tuple[str, float]]]:
    """Databank workbook → {series_key: [(quarter_start, value)]}."""
    html = requests.get(URLS["hhdc_index"], timeout=(10, 60)).text
    vintages = re.findall(r"hhd_c_report_(\d{4})q(\d)\.xlsx", html)
    if not vintages:
        raise NyFedResearchError("hhdc: no report files in databank.html")
    year, q = max((int(y), int(n)) for y, n in vintages)
    url = f"{BASE_WWW}/medialibrary/interactives/householdcredit/data/xls/hhd_c_report_{year}q{q}.xlsx"
    if url not in _bytes_cache:
        r = requests.get(url, timeout=(10, 120))
        if r.status_code != 200 or r.content[:2] != b"PK":
            raise NyFedResearchError(f"hhdc: HTTP {r.status_code} / not XLSX ({year}q{q})")
        _bytes_cache[url] = r.content

    import openpyxl

    # data_only=True: 'Total' cells are formulas — without it we store '=SUM(...)'
    wb = openpyxl.load_workbook(io.BytesIO(_bytes_cache[url]), read_only=True, data_only=True)

    def _sheet_rows(name: str) -> list[list]:
        if name not in wb.sheetnames:
            raise NyFedResearchError(f"hhdc: sheet '{name}' missing (sheets moved?)")
        return [list(r) for r in wb[name].iter_rows(values_only=True)]

    def _by_header(rows: list[list], must_contain: str) -> tuple[int, dict[str, int]]:
        """Locate the header row + column map BY NAME (pages reorder columns
        between vintages — Page 12 is MORTGAGE-first, Page 14 is AUTO-first)."""
        for i, row in enumerate(rows[:10]):
            labels = [str(c).strip().upper() if c is not None else "" for c in row]
            if must_contain in labels:
                return i, {lab: j for j, lab in enumerate(labels) if lab}
        raise NyFedResearchError(f"hhdc: header with '{must_contain}' not found")

    out: dict[str, list[tuple[str, float]]] = {
        k: [] for k in (
            "HHDC_TOTAL_DEBT", "HHDC_ORIG_MORTGAGE", "HHDC_ORIG_AUTO",
            "HHDC_DQ90_FLOW_MORTGAGE", "HHDC_DQ90_FLOW_AUTO",
            "HHDC_DQ90_FLOW_CC", "HHDC_DQ90_FLOW_STUDENT",
        )
    }

    # Page 3: balances — TRILLIONS → bn_usd (×1000, the single conversion site)
    rows = _sheet_rows("Page 3 Data")
    hi, cols = _by_header(rows, "MORTGAGE")
    for row in rows[hi + 1:]:
        ts = _quarter_start(row[0]) if row and _is_quarter_label(row[0]) else None
        v = _f(row[cols["TOTAL"]]) if ts else None
        if ts and v is not None:
            out["HHDC_TOTAL_DEBT"].append((ts, round(v * 1000, 1)))

    # Pages 6/8: originations ($bn) — 'Total' column
    for page, key in (("Page 6 Data", "HHDC_ORIG_MORTGAGE"), ("Page 8 Data", "HHDC_ORIG_AUTO")):
        rows = _sheet_rows(page)
        hi, cols = _by_header(rows, "TOTAL")
        for row in rows[hi + 1:]:
            ts = _quarter_start(row[0]) if row and _is_quarter_label(row[0]) else None
            v = _f(row[cols["TOTAL"]]) if ts else None
            if ts and v is not None:
                out[key].append((ts, round(v, 2)))

    # Page 14: 90+ delinquency FLOW rates (percent) by loan type
    rows = _sheet_rows("Page 14 Data")
    hi, cols = _by_header(rows, "AUTO")
    flow_map = {  # labels exactly as the sheet prints them (live-verified)
        "HHDC_DQ90_FLOW_MORTGAGE": "MORTGAGE",
        "HHDC_DQ90_FLOW_AUTO": "AUTO",
        "HHDC_DQ90_FLOW_CC": "CC",
        "HHDC_DQ90_FLOW_STUDENT": "STUDENT LOAN",
    }
    for row in rows[hi + 1:]:
        ts = _quarter_start(row[0]) if row and _is_quarter_label(row[0]) else None
        if not ts:
            continue
        for key, label in flow_map.items():
            v = _f(row[cols[label]]) if label in cols else None
            if v is not None:
                out[key].append((ts, round(v, 3)))
    wb.close()

    if not out["HHDC_TOTAL_DEBT"]:
        raise NyFedResearchError("hhdc: no balance rows parsed")
    return out


# --- SCE (monthly expectations) ----------------------------------------------------


# (series_key, sheet, label-prefix) — labels live-verified 2026-09-19
_SCE_SPECS = [
    ("SCE_INFL_1Y", "Inflation expectations", "Median one-year ahead expected inflation"),
    ("SCE_INFL_3Y", "Inflation expectations", "Median three-year ahead expected inflation"),
    ("SCE_INFL_5Y", "Five-year ahead Infl Exp", "Median five-year ahead expected inflation"),
    ("SCE_EARN", "Earnings growth", "Median expected earnings growth"),
    ("SCE_JOBLOSS", "Job separation expectation", "Mean probability of losing a job"),
    ("SCE_JOBFIND", "Job finding expectations", "Mean probability of finding a job"),
]


def _sce() -> dict[str, list[tuple[str, float]]]:
    """Core-module workbook → series. Columns are mapped BY LABEL (row idx 3):
    job-loss and quit probabilities share a sheet — a positional read swaps
    them (a claim-drift class defect)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(_download("sce")), read_only=True, data_only=True)
    out: dict[str, list[tuple[str, float]]] = {}
    for key, sheet, label in _SCE_SPECS:
        ws = wb[sheet]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        col = next(
            (j for j, c in enumerate(rows[3]) if isinstance(c, str) and c.strip().startswith(label)),
            None,
        )
        if col is None:
            raise NyFedResearchError(f"sce: label {label!r} not found in sheet {sheet!r}")
        vals = []
        for row in rows[4:]:
            ym = row and row[0]
            v = _f(row[col]) if col < len(row) else None
            if isinstance(ym, (int, float)) and int(ym) > 190000 and v is not None:
                ym = int(ym)
                vals.append((f"{ym // 100}-{ym % 100:02d}-01", round(v, 4)))
        if not vals:
            raise NyFedResearchError(f"sce: no rows for {key}")
        out[key] = vals
    wb.close()
    return out


# --- GSCPI (monthly index, vintage matrix) ------------------------------------------


def _gscpi() -> dict[str, list[tuple[str, float]]]:
    """Interactive CSV → current-vintage series (LAST column)."""
    text = _download("gscpi").decode("utf-8", errors="replace")
    rdr = csv.reader(io.StringIO(text))
    header = next(rdr)
    cur = len(header) - 1  # current vintage = last column
    vals = []
    for row in rdr:
        if len(row) <= cur or not row[0].strip():
            continue
        v = _f(row[cur])
        if v is not None:
            # source labels are month-END ('31-Aug-2026') — normalize to month
            # START: the report_month convention every other monthly series
            # uses (audit 2026-09-20: GSCPI was the sole day-28..31 violator)
            vals.append((_month_start(_dd_mon_yyyy(row[0])), round(v, 4)))
    if not vals:
        raise NyFedResearchError("gscpi: no rows in current vintage column")
    return {"GSCPI": vals}


# --- Empire State (monthly diffusion) ------------------------------------------------


def _empire() -> dict[str, list[tuple[str, float]]]:
    """SA diffusion CSV. surveyDate is the survey month's END — normalized to
    month START (report_month registry convention)."""
    text = _download("empire").decode("utf-8", errors="replace")
    rdr = csv.DictReader(io.StringIO(text))
    colmap = {
        "ESMS_HEADLINE": "GACDISA",
        "ESMS_NEW_ORDERS": "NOCDISA",
        "ESMS_SHIPMENTS": "SHCDISA",
        "ESMS_PRICES_PAID": "PPCDISA",
        "ESMS_PRICES_RECEIVED": "PRCDISA",
        "ESMS_EMPLOYMENT": "NECDISA",
        "ESMS_AVG_WORKWEEK": "AWCDISA",
    }
    out: dict[str, list[tuple[str, float]]] = {k: [] for k in colmap}
    for row in rdr:
        d = (row.get("surveyDate") or "").strip()
        if not re.match(r"\d{4}-\d{2}-\d{2}", d):
            continue
        ts = _month_start(d)
        for key, col in colmap.items():
            v = _f(row.get(col))
            if v is not None:
                out[key].append((ts, round(v, 2)))
    if not out["ESMS_HEADLINE"]:
        raise NyFedResearchError("empire: no diffusion rows")
    return out


# --- ACM term premia (daily, legacy xls) --------------------------------------------


def _acm() -> dict[str, list[tuple[str, float]]]:
    """ACMTermPremium.xls 'ACM Daily' → 1y/2y/5y/10y term premia (percent).
    Legacy OLE2 → xlrd (openpyxl cannot read it); all cells are TEXT."""
    import xlrd

    wb = xlrd.open_workbook(file_contents=_download("acm"))
    ws = wb.sheet_by_name("ACM Daily")
    header = [str(ws.cell_value(0, c)).strip() for c in range(ws.ncols)]
    cols = {
        "ACMTP01": header.index("ACMTP01"),
        "ACMTP02": header.index("ACMTP02"),
        "ACMTP05": header.index("ACMTP05"),
        "ACMTP10": header.index("ACMTP10"),
    }
    out = {k: [] for k in cols}
    for r in range(1, ws.nrows):
        d = str(ws.cell_value(r, 0)).strip()
        if not d:
            continue
        try:
            ts = _dd_mon_yyyy(d)
        except ValueError:
            continue
        for key, c in cols.items():
            v = _f(ws.cell_value(r, c))
            if v is not None:
                out[key].append((ts, round(v, 4)))
    if not out["ACMTP10"]:
        raise NyFedResearchError("acm: no rows")
    return out


# --- HPW labor-market tightness (monthly z-score) ------------------------------------


def _hpw() -> dict[str, list[tuple[str, float]]]:
    """HPW index — z-SCORE units (mean 0 / sd 1, NOT percent). The trailing
    month row is an empty placeholder: dropped by the value-not-None filter."""
    text = _download("hpw").decode("utf-8", errors="replace")
    rdr = csv.DictReader(io.StringIO(text))
    vals = []
    for row in rdr:
        v = _f(row.get("HPW"))
        if v is not None and (row.get("date") or "").strip():
            vals.append((_m_dy(row["date"]), round(v, 5)))
    if not vals:
        raise NyFedResearchError("hpw: no rows")
    return {"HPW_IDX_M": vals}


# --- Laubach-Williams r-star (quarterly, whole-history revisions) ---------------------


def _lw() -> dict[str, list[tuple[str, float]]]:
    """LW one-sided rstar + output gap, HLW US r-star. The first 'rstar' column
    whose GROUP label (row above the header) says One-Sided is taken — the
    two-sided block repeats the same names (terminal rows coincide)."""
    import openpyxl

    out: dict[str, list[tuple[str, float]]] = {"LW_RSTAR": [], "LW_GAP": [], "HLW_RSTAR": []}

    def _estimates(path_key: str, sheet: str) -> list[list]:
        wb = openpyxl.load_workbook(io.BytesIO(_download(path_key)), read_only=True, data_only=True)
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        wb.close()
        return rows

    def _header_idx(rows: list[list]) -> int | None:
        return next(
            (i for i, r in enumerate(rows[:10]) if r and str(r[0] or "").strip() == "Date"),
            None,
        )

    # --- LW file: 'data' sheet, one-sided cols C(rstar)..F(gap) ---
    rows = _estimates("lw", "data")
    hi = _header_idx(rows)
    if hi is None:
        raise NyFedResearchError("lw: 'Date' header not found")
    groups = rows[hi - 1] if hi >= 1 else []
    names = rows[hi]
    rstar_col = next(
        (
            j
            for j, c in enumerate(names)
            if str(c or "").strip().lower() == "rstar"
            and "one" in str(groups[j] if j < len(groups) else "").lower()
        ),
        None,
    )
    if rstar_col is None:  # group row may be merged/blank — fall back to FIRST rstar col
        rstar_col = next(
            (j for j, c in enumerate(names) if str(c or "").strip().lower() == "rstar"), None
        )
    if rstar_col is None:
        raise NyFedResearchError("lw: rstar column not found")
    gap_col = next(
        (
            j
            for j, c in enumerate(names)
            if j > rstar_col and "gap" in str(c or "").strip().lower()
        ),
        rstar_col + 3,
    )
    for row in rows[hi + 1:]:
        if not row or not isinstance(row[0], datetime):
            continue
        ts = f"{row[0].year}-{(row[0].month - 1) // 3 * 3 + 1:02d}-01"
        v = _f(row[rstar_col]) if rstar_col < len(row) else None
        if v is not None:
            out["LW_RSTAR"].append((ts, round(v, 4)))
        g = _f(row[gap_col]) if gap_col < len(row) else None
        if g is not None:
            out["LW_GAP"].append((ts, round(g, 4)))

    # --- HLW file: 'HLW Estimates', US r* = the 'US' column whose GROUP label
    # (row above the header) says 'Natural Rate' — the group label sits directly
    # above the US column (live-caught: grp_col+1 landed on CANADA's 1.76)
    rows = _estimates("hlw", "HLW Estimates")
    hi = _header_idx(rows)
    groups = rows[hi - 1] if hi >= 1 else []
    rstar_col = next(
        (
            j
            for j, c in enumerate(rows[hi])
            if str(c or "").strip() == "US"
            and "natural rate" in str(groups[j] if j < len(groups) else "").lower()
        ),
        None,
    )
    if rstar_col is None:
        raise NyFedResearchError("hlw: US r* column (Natural Rate group) not found")
    for row in rows[hi + 1:]:
        if not row or not isinstance(row[0], datetime):
            continue
        v = _f(row[rstar_col]) if rstar_col < len(row) else None
        if v is not None:
            ts = f"{row[0].year}-{(row[0].month - 1) // 3 * 3 + 1:02d}-01"
            out["HLW_RSTAR"].append((ts, round(v, 4)))

    if not out["LW_RSTAR"] or not out["HLW_RSTAR"]:
        raise NyFedResearchError("lw: no estimate rows parsed")
    return out


# --- MCT core inflation (monthly) ----------------------------------------------------


def _mct() -> dict[str, list[tuple[str, float]]]:
    """Chart CSV → MCT point + 68% band. The four 'MCT' columns are
    band-lo | point | band-hi | vs-2017-19-avg — enforced lo<=point<=hi at
    parse (a schema-drift tripwire, semantics verified numerically 2026-09-19)."""
    text = _download("mct").decode("utf-8", errors="replace")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if len(lines) < 6:
        raise NyFedResearchError("mct: file too short")
    out: dict[str, list[tuple[str, float]]] = {"MCT": [], "MCT_BAND_LO": [], "MCT_BAND_HI": []}
    for ln in lines[4:]:  # 4 header rows; the leading unnamed col shifts dates to parts[1]
        parts = ln.split(",")
        if len(parts) < 5 or not re.match(r"\d{1,2}/\d{1,2}/\d{4}", parts[1].strip()):
            continue
        ts = _m_dy(parts[1])
        lo, pt, hi = _f(parts[2]), _f(parts[3]), _f(parts[4])
        if None in (lo, pt, hi):
            continue
        if not (lo <= pt <= hi):
            raise NyFedResearchError(
                f"mct: band invariant broken {ts}: lo={lo} pt={pt} hi={hi} (schema drift?)"
            )
        out["MCT_BAND_LO"].append((ts, round(lo, 3)))
        out["MCT"].append((ts, round(pt, 3)))
        out["MCT_BAND_HI"].append((ts, round(hi, 3)))
    if not out["MCT"]:
        raise NyFedResearchError("mct: no rows")
    return out


FAMILY_PARSERS = {
    "hhdc": _hhdc,
    "sce": _sce,
    "gscpi": _gscpi,
    "empire": _empire,
    "acm": _acm,
    "hpw": _hpw,
    "lw": _lw,
    "mct": _mct,
}

# series key (sans NYFED:) → family. Keys must match the parser outputs above.
SERIES_FAMILY = {
    "HHDC_TOTAL_DEBT": "hhdc", "HHDC_ORIG_MORTGAGE": "hhdc", "HHDC_ORIG_AUTO": "hhdc",
    "HHDC_DQ90_FLOW_MORTGAGE": "hhdc", "HHDC_DQ90_FLOW_AUTO": "hhdc",
    "HHDC_DQ90_FLOW_CC": "hhdc", "HHDC_DQ90_FLOW_STUDENT": "hhdc",
    "SCE_INFL_1Y": "sce", "SCE_INFL_3Y": "sce", "SCE_INFL_5Y": "sce", "SCE_EARN": "sce",
    "SCE_JOBLOSS": "sce", "SCE_JOBFIND": "sce",
    "GSCPI": "gscpi",
    "ESMS_HEADLINE": "empire", "ESMS_NEW_ORDERS": "empire", "ESMS_SHIPMENTS": "empire",
    "ESMS_PRICES_PAID": "empire", "ESMS_PRICES_RECEIVED": "empire",
    "ESMS_EMPLOYMENT": "empire", "ESMS_AVG_WORKWEEK": "empire",
    "ACMTP01": "acm", "ACMTP02": "acm", "ACMTP05": "acm", "ACMTP10": "acm",
    "HPW_IDX_M": "hpw",
    "LW_RSTAR": "lw", "LW_GAP": "lw", "HLW_RSTAR": "lw",
    "MCT": "mct", "MCT_BAND_LO": "mct", "MCT_BAND_HI": "mct",
}

# window floors: a 10-day default would starve monthly/quarterly families —
# the window must always reach at least the latest observation
FLOOR_DAYS = {
    "hhdc": 420, "sce": 120, "gscpi": 150, "empire": 150, "hpw": 150,
    "lw": 550, "mct": 150, "acm": 10,
}

_parsed_cache: dict[str, dict[str, list[tuple[str, float]]]] = {}


def _family_rows(family: str) -> dict[str, list[tuple[str, float]]]:
    if family not in _parsed_cache:
        _parsed_cache[family] = FAMILY_PARSERS[family]()
    return _parsed_cache[family]


def knows(series_key: str) -> bool:
    return series_key in SERIES_FAMILY


def fetch_latest(series_id: str) -> dict:
    key = series_id.split(":", 1)[1] if ":" in series_id else series_id
    family = SERIES_FAMILY.get(key)
    if family is None:
        raise NyFedResearchError(f"nyfedresearch: unrouted series {series_id}")
    rows = _family_rows(family).get(key) or []
    if not rows:
        raise NyFedResearchError(f"nyfedresearch: {key}: no rows")
    ts, v = max(rows, key=lambda x: x[0])
    return {"ts": ts, "value": v}


def fetch_window(series_id: str, days: int = 10, *, today: str | None = None) -> list[dict]:
    """`today` (ISO) is injectable so tests pin the cutoff instead of racing
    the wall clock (audit 2026-09-20: the floor tests were time bombs —
    frozen 2026 quarters vs a moving 'now' would fail spuriously in 2027)."""
    key = series_id.split(":", 1)[1] if ":" in series_id else series_id
    family = SERIES_FAMILY.get(key)
    if family is None:
        raise NyFedResearchError(f"nyfedresearch: unrouted series {series_id}")
    from datetime import UTC, timedelta

    now = datetime.now(UTC) if today is None else datetime.fromisoformat(today + "T00:00:00+00:00")
    cutoff = (now - timedelta(days=max(days, FLOOR_DAYS[family]))).date().isoformat()
    return [{"ts": ts, "value": v} for ts, v in _family_rows(family).get(key, []) if ts >= cutoff]


def save_history(conn, *, verbose: bool = True) -> int:
    """Full-history ingest of every REGISTERED nyfedresearch series (FK-safe:
    unregistered keys are skipped — the sep.py lesson). Revisions land via
    db.apply_realtime_revisions (vintage snapshot + realtime update)."""
    from .. import db

    # registry ids via the qa loader (lazy import avoids the fetchers→qa cycle
    # at module import time; harvest imports this module eagerly)
    from ..qa.verify_sources import load_registry

    reg_ids = {e["series_id"] for e in load_registry()}
    total = 0
    for key in SERIES_FAMILY:
        sid = f"NYFED:{key}"
        if sid not in reg_ids:
            continue
        rows = _family_rows(SERIES_FAMILY[key]).get(key) or []
        payload = [(sid, ts, v, "NYFED") for ts, v in rows]
        n_new = db.insert_observations(conn, payload)
        n_rev = db.apply_realtime_revisions(conn, payload)
        total += n_new + n_rev
        if verbose and rows:
            print(f"  {sid:28s} {len(rows):4d} obs (new {n_new}, revised {n_rev}) — latest {rows[-1]}")
    conn.commit()
    return total
