"""flows_extra.py — flow fetchers: Farside BTC/ETH ETF flows (per-issuer),
TIC SLT foreign Treasury holdings, PBoC/SAFE reserve assets, LME copper
stocks, and LBMA London vault data (gold + silver).

The targets are bot-gated or ship unusual formats (Farside HTML tables with
'DD MMM YYYY' dates, Treasury TIC HTML, SAFE/LME/LBMA XLSX archives), so each
parser is tailored to a verified response shape.

TRAP (F-K6 sibling, ISSUE D-021): never derive "latest" from DOM position —
Farside renders date rows OLDEST-FIRST, so trs[0] was 19 days stale while the
fetch exited 0. Always pick max() over the parsed date values.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime

import requests
from curl_cffi import requests as creq

UA = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128.0 Safari/537.36"
}


def _signed(open_p: str, raw: str, close_p: str) -> float:
    """Parenthesized '(1,234.5)' = outflow. Parens must be detected in the RAW
    match — after comma-stripping the search string no longer matches
    (F-K6)."""
    val = float(raw.replace(",", "").rstrip("."))
    return -val if (open_p and close_p) else val


_FARSIDE_NUM = re.compile(
    r'<span class="(?:tabletext|redFont|greenFont)">(\(?)([\d.,]+)(\))?</span>'
)


def _fetch_farside(path: str) -> dict:
    """Full window parse of a Farside ETF-flow table.

    Returns {'rows': [...], 'issuers': [...], 'latest': {...}, 'cumulative':
    {...}} where each row is {'ts', 'date_iso', 'net_flow_musd', 'issuers':
    {TICKER: signed $M}} and 'latest' is the max-DATE row (never the DOM-first
    row). Rows whose only numeric cell is a bare '0.0' are intraday skeleton
    placeholders and are skipped — landing one would write a fake zero-flow
    day.
    """
    r = requests.get(f"https://farside.co.uk/{path}/", headers=UA, timeout=(10, 30))
    if r.status_code != 200:
        raise RuntimeError(f"Farside {path}: HTTP {r.status_code}")
    text = r.text

    # issuer names: the thead run between 'Total' and 'Fee' (both stripped of
    # tags/&nbsp). Live-verified order matches the data cells, Total excluded.
    issuers: list[str] = []
    thead = re.search(r"<thead>(.*?)</thead>", text, re.S)
    if thead:
        th = [
            re.sub(r"&nbsp;|\s+", "", re.sub(r"<[^>]+>", "", c)).strip()
            for c in re.findall(r"<th[^>]*>(.*?)</th>", thead.group(1), re.S)
        ]
        if "Total" in th and "Fee" in th:
            issuers = [t for t in th[th.index("Total") + 1 : th.index("Fee")] if t]
    if not issuers:
        raise RuntimeError(f"Farside {path}: issuer header row not parsed")

    # date rows: 'DD MMM YYYY' key + the numeric cells of that <tr>
    trs = re.findall(
        r'<tr[^>]*>\s*<td><span class="tabletext">(\d{1,2}\s+\w{3}\s+\d{4})</span></td>(.*?)</tr>',
        text,
        re.S,
    )
    if not trs:
        raise RuntimeError(f"Farside {path}: no date rows in tbody")
    rows = []
    for date, rest in trs:
        nums = _FARSIDE_NUM.findall(rest)
        try:
            iso = datetime.strptime(date, "%d %b %Y").date().isoformat()
        except ValueError:
            continue  # non-calendar date would poison the max-date pick
        signed = [_signed(o, v, c) for o, v, c in nums]
        # Cell-count contract (live-verified): a COMPLETE day renders
        # len(issuers)+1 cells (issuer flows + Total). Fewer = intraday
        # partial: issuer-only rows are landed per-issuer but their aggregate
        # is NOT guessed (misreading the BTC-mini cell as 'Total' would write
        # a wrong flows_daily net); sub-issuer rows are skipped and the daily
        # whole-window upsert self-heals them on the next run.
        if len(signed) == len(issuers) + 1:
            per_issuer, total = dict(zip(issuers, signed[:-1], strict=True)), signed[-1]
        elif len(signed) == len(issuers):
            per_issuer, total = dict(zip(issuers, signed, strict=True)), None
        else:
            continue
        rows.append(
            {"ts": date, "date_iso": iso, "net_flow_musd": total, "issuers": per_issuer}
        )
    if not rows:
        raise RuntimeError(f"Farside {path}: no filled date rows")

    # footer: since-inception cumulative per issuer, then net total
    cumulative: dict[str, float] = {}
    m = re.search(
        r'<tr[^>]*>\s*<td[^>]*>\s*<span class="tabletext">\s*Total\s*</span>(.*?)</tr>',
        text,
        re.S,
    )
    if m:
        cums = _FARSIDE_NUM.findall(m.group(1))
        if len(cums) == len(issuers) + 1:
            cumulative = dict(
                zip(issuers, [_signed(o, v, c) for o, v, c in cums[:-1]], strict=False)
            )
            cumulative["Total"] = _signed(*cums[-1])

    latest = max(rows, key=lambda x: x["date_iso"])
    return {"rows": rows, "issuers": issuers, "latest": latest, "cumulative": cumulative}


def fetch_farside_btc() -> dict:
    return _fetch_farside("btc")


def fetch_farside_eth() -> dict:
    return _fetch_farside("eth")


# rows extracted per country from the SLT table (label -> flows_periodic kind)
_TIC_ROWS = {
    "China, Mainland": "china",
    "Belgium": "belgium",
    "Japan": "japan",
    "Cayman Islands": "cayman",
    "United Kingdom": "united_kingdom",
    "Grand Total": "grand_total",
    "Of Which: Foreign Official Treasury Bills": "official_bills",
}


def fetch_tic_slt5() -> dict:
    """Foreign holdings of US Treasuries from the TIC SLT table5 (monthly).

    The table displays a SINGLE populated column (the latest period) — the
    dozen 'YYYY-mm' strings elsewhere in the HTML are a month picker, not
    filled data columns, so there is no same-page history to read. Belgium =
    the classic Euroclear 'stealth China' custody proxy; Grand Total pairs
    with FISCAL:DEBT_TOTAL for the term-premium absorption story.
    """
    s = creq.Session(impersonate="chrome")
    r = s.get(
        "https://ticdata.treasury.gov/resource-center/data-chart-center/tic/Documents/"
        "slt_table5.html",
        timeout=(10, 30),
    )
    if r.status_code != 200:
        raise RuntimeError(f"TIC: HTTP {r.status_code}")
    cells = [c.strip() for c in re.sub(r"<[^>]+>", "|", r.text).split("|")]
    periods, run = [], []
    for c in cells:
        if not c:
            continue
        if re.fullmatch(r"20\d\d-\d\d", c):
            run.append(c)
        else:
            if len(run) >= 3:
                periods.extend(run)
            run = []
    if len(run) >= 3:
        periods.extend(run)
    values: dict[str, float] = {}
    for label, key in _TIC_ROWS.items():
        try:
            i = cells.index(label)
        except ValueError:
            continue  # row absent this month → kind simply not refreshed
        for c in cells[i + 1 :]:
            if re.fullmatch(r"[\d,.]+", c or ""):
                values[key] = float(c.replace(",", ""))
                break
            if c:  # first non-empty non-number cell ends the row
                break
    if "china" not in values:
        raise RuntimeError("TIC: China, Mainland row not found")
    return {
        "ts": periods[0] if periods else "latest",
        "values": values,  # {key: $B}, china guaranteed
    }


def fetch_pboc_gold() -> dict:
    """PBoC gold + reserve-asset composition from the SAFE monthly XLSX.

    Layout (verified): row 3 date headers 'YYYY.MM' at ODD columns; each month
    spans a USD/SDR column PAIR; the gold tonnage text 'N万盎司' is duplicated
    across both columns of its pair. Reading tonnage at the odd (USD) columns
    only is what binds each value to the correct month — scanning from the
    right lands on the unlabelled SDR sub-column and, with a now() fallback,
    mislabels the datum to the RUN month (ISSUE D-021).

    Returns {'months': [{ts, wan_oz, tonnes}], 'latest': {...},
    'fx_reserves_usd_yi', 'gold_value_usd_yi', 'total_reserves_usd_yi',
    'gold_share_pct'} — months cover the current sheet year (Jan→latest).
    """
    from openpyxl import load_workbook

    from ..units import wan_oz_to_tonnes

    s = creq.Session(impersonate="chrome")
    r = s.get("https://www.safe.gov.cn/en/2021/0203/2045.html", timeout=(10, 30))
    if r.status_code != 200:
        raise RuntimeError(f"SAFE: HTTP {r.status_code}")
    links = re.findall(r'href="(/en/file/file/[^"]+\.xlsx)"', r.text)
    if not links:
        raise RuntimeError("SAFE: XLSX link not found")
    r2 = s.get(f"https://www.safe.gov.cn{links[0]}", timeout=(10, 60))
    if r2.status_code != 200 or r2.content[:2] != b"PK":
        raise RuntimeError(f"SAFE XLSX: HTTP {r2.status_code}")
    wb = load_workbook(io.BytesIO(r2.content), read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))

    # month -> USD column index (odd cols) from the row-3 headers
    month_cols: list[tuple[str, int]] = []
    for i, c in enumerate(rows[3] if len(rows) > 3 else []):
        m = re.match(r"(\d{4})\.(\d{2})", str(c).strip() if c else "")
        if m and i % 2 == 1:
            month_cols.append((f"{m.group(1)}-{m.group(2)}", i))

    def _latest_usd(row_vals: list[str]) -> float | None:
        """Last filled ODD-column numeric on a 亿美元 row (SDR cells skipped)."""
        found = None
        for _ts, col in month_cols:
            if col < len(row_vals) and re.fullmatch(r"[-\d.,]+", row_vals[col] or ""):
                found = float(row_vals[col].replace(",", ""))
        return found

    fx = gold_val = total = None
    tonnage_rows: list[tuple[str, float]] = []
    for row in rows:
        vals = [str(c).strip() if c is not None else "" for c in row]
        label = vals[0] if vals else ""
        if "外汇储备" in label:
            fx = _latest_usd(vals)
        elif "黄金" in label:
            gold_val = _latest_usd(vals)
        elif "合计" in label:
            total = _latest_usd(vals)
        elif not label and any("万盎司" in v for v in vals):
            # tonnage row: read the USD (odd) columns of each pair
            for ts, col in month_cols:
                m = re.search(r"([\d.,]+)\s*万盎司", vals[col] if col < len(vals) else "")
                if m:
                    tonnage_rows.append((ts, float(m.group(1).replace(",", ""))))
    if not tonnage_rows:
        raise RuntimeError("SAFE: gold volume row (万盎司) not found")

    months = [
        {"ts": ts, "wan_oz": w, "tonnes": round(wan_oz_to_tonnes(w), 1)}
        for ts, w in tonnage_rows
    ]
    latest = months[-1]
    gold_share = round(gold_val / total * 100, 2) if gold_val and total else None
    return {
        "months": months,
        "latest": latest,
        # legacy keys kept: f2 stored ts/wan_oz/tonnes at the top level
        "ts": latest["ts"],
        "wan_oz": latest["wan_oz"],
        "tonnes": latest["tonnes"],
        "fx_reserves_usd_yi": fx,
        "gold_value_usd_yi": gold_val,
        "total_reserves_usd_yi": total,
        "gold_share_pct": gold_share,
    }


_LME_MONTHS = (
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
)


def fetch_lme_stocks(year: int, month: int, session=None) -> list[dict]:
    """LME daily closing stocks (CA = copper column), unit tonnes.

    Monthly XLSX archives follow a predictable filename pattern,
    stocks-{month}-{year}.xlsx (archive reaches back to 2018). Two header
    eras exist ('Stock Date' vs 'BusinessDate'), so the header row is
    located dynamically. Cloudflare requires curl_cffi chrome impersonation,
    and the session must be SHARED across months: a fresh session per file
    fires many consecutive distinct TLS fingerprints and gets blocked.
    Returns [{'ts': 'YYYY-MM-DD', 'copper_tonnes': float}] daily rows; []
    when the file is not published yet or the format is unrecognized.
    """
    from openpyxl import load_workbook

    url = (
        "https://www.lme.com/-/media/files/data/reports-and-data/"
        f"warehouse-and-stock-reports/stocks-summary/"
        f"stocks-{_LME_MONTHS[month - 1]}-{year}.xlsx"
    )
    s = session if session is not None else creq.Session(impersonate="chrome")
    r = s.get(url, timeout=(10, 60))
    # Unpublished months are a soft-404: HTTP 200 with an HTML body.
    # Genuine XLSX starts with the PK magic bytes.
    if r.status_code != 200 or r.content[:2] != b"PK":
        return []
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    # locate the header row dynamically: 'BusinessDate'/'Stock Date' plus 'CA'
    hdr_i = date_col = ca_col = None
    for i, row in enumerate(rows[:20]):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "BusinessDate" in cells:
            hdr_i, date_col = i, cells.index("BusinessDate")
        elif "Stock Date" in cells:
            hdr_i, date_col = i, cells.index("Stock Date")
        if hdr_i is not None:
            if "CA" in cells:
                ca_col = cells.index("CA")
                break
            hdr_i = None  # header row without a CA column -> keep looking
    if hdr_i is None or ca_col is None:
        return []
    out = []
    for row in rows[hdr_i + 1 :]:
        if len(row) <= ca_col:
            continue
        d, v = row[date_col], row[ca_col]
        if d is None or v is None:
            continue
        ts = d.date().isoformat() if hasattr(d, "date") else str(d)[:10]
        try:
            out.append({"ts": ts, "copper_tonnes": float(v)})
        except (TypeError, ValueError):
            continue
    return out


_LBMA_MONTHS = {
    "january": "01",
    "february": "02",
    "march": "03",
    "april": "04",
    "may": "05",
    "june": "06",
    "july": "07",
    "august": "08",
    "september": "09",
    "october": "10",
    "november": "11",
    "december": "12",
}


_LME_OW_MONTHS = {m: f"{i:02d}" for i, m in enumerate(
    ("january", "february", "march", "april", "may", "june", "july",
     "august", "september", "october", "november", "december"), start=1
)}


def fetch_lme_offwarrant(session=None) -> list[dict]:
    """Off-warrant copper stocks from the LME monthly archive (tonnes).

    Context (Notice 25/054, verified 2026-09-10): since 2025-04 the DAILY
    off-warrant reports are T+1 paid ($1,200/yr licensing) or T+3 free via an
    LME.com LOGIN; the anonymous monthly archive here is FROZEN at 2025-02.
    So this is a baseline/backfill series — the off-warrant share of total
    stocks before reporting moved behind login — not a live feed. The page
    lists every monthly file, so if the LME ever re-publishes anonymously,
    the next harvest picks it up with zero changes. Returns
    [{'ts': 'YYYY-MM', 'cu_tonnes': float, 'regions': {..}}], newest first.
    """
    from openpyxl import load_workbook

    s = session if session is not None else creq.Session(impersonate="chrome")
    r = s.get(
        "https://www.lme.com/market-data/reports-and-data/warehouse-and-stocks-reports/"
        "off-warrant-stock-reporting",
        timeout=(10, 30),
    )
    if r.status_code != 200:
        raise RuntimeError(f"LME off-warrant: HTTP {r.status_code}")
    links = sorted(set(re.findall(
        r'href="(/-/media/files/data/reports-and-data/warehouse-and-stock-reports/'
        r'off-warrant-stock-reporting/[^"]+\.xlsx)"', r.text
    )))
    out = []
    for href in links:
        m = re.search(r"-([a-z]+)-(\d{4})\.xlsx$", href, re.I)
        mm = _LME_OW_MONTHS.get(m.group(1).lower()) if m else None
        if not mm:
            continue
        ts = f"{m.group(2)}-{mm}"
        r2 = s.get(f"https://www.lme.com{href}", timeout=(10, 60))
        if r2.status_code != 200 or r2.content[:2] != b"PK":
            continue
        wb = load_workbook(io.BytesIO(r2.content), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        # header: 'Region|Location|AA|AL|CU|NA|NI|PB|SN|ZN'
        hdr_i = cu_col = None
        for i, row in enumerate(rows[:10]):
            cells = [str(c).strip().upper() if c is not None else "" for c in row]
            if "LOCATION" in cells and "CU" in cells:
                hdr_i, cu_col = i, cells.index("CU")
                break
        if hdr_i is None:
            continue
        regions: dict[str, float] = {}
        total = 0.0
        for row in rows[hdr_i + 1 :]:
            cells = [str(c).strip() if c is not None else "" for c in row]
            if len(cells) <= cu_col or not cells[0]:
                continue
            label = cells[0].upper()
            if not label.startswith("TOTAL"):
                continue
            region = label.replace("TOTAL", "").strip().title()
            if not region:
                continue  # the bare grand-TOTAL row would double-count
            v = _num_or(cells[cu_col])
            regions[region] = v
            total += v
        if total <= 0:
            continue
        out.append({"ts": ts, "cu_tonnes": total, "regions": regions})
    out.sort(key=lambda x: x["ts"], reverse=True)
    return out


def _num_or(v) -> float:
    """LME cells: numbers, '/' (not held), or empty — '/' must be 0."""
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def fetch_lbma_vault() -> dict:
    """LBMA London vault holdings (gold + silver) — FULL history in one XLSX.

    The single file carries every month since 2016-07 (verified: 125 rows),
    so the daily harvest doubles as the backfill. Gold is the physical-gold
    map's London leg next to PBoC tonnage and GLD shares; the old fallback
    URL pattern moved to /downloads/ when the CDN was restructured
    (media/lbv-*.xlsx is a guaranteed 403 today).
    """
    from openpyxl import load_workbook

    from ..units import koz_to_tonnes

    s = creq.Session(impersonate="chrome")
    r = s.get("https://www.lbma.org.uk/prices-and-data/london-vault-data", timeout=(10, 30))
    links = (
        re.findall(r'href="(https://cdn\.lbma\.org\.uk/[^"]+\.xlsx)"', r.text)
        if r.status_code == 200
        else []
    )
    if not links:
        now = datetime.now(UTC)
        py, pm = (now.year, now.month - 1) if now.month > 1 else (now.year - 1, 12)
        name = next(n for n, mm in _LBMA_MONTHS.items() if mm == f"{pm:02d}")
        links = [f"https://cdn.lbma.org.uk/downloads/LBMA-London-Vault-Holdings-Data-{name}-{py}.xlsx"]
    r2 = s.get(links[0], timeout=(10, 60))
    if r2.status_code != 200 or r2.content[:2] != b"PK":
        raise RuntimeError(f"LBMA: HTTP {r2.status_code}")
    wb = load_workbook(io.BytesIO(r2.content), read_only=True, data_only=True)
    # Layout (verified 2026-09): row 1 headers ('Month End|Gold|Silver'), row
    # 2 units ("Troy Ounces ('000s)"), row 3+ data with 'YYYY-MM' string
    # labels (newer) or datetime cells (2016-2023 era). Newest-first.
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    gold_col = silver_col = None
    for row in rows[:3]:
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if "gold" in cells and "silver" in cells:
            gold_col, silver_col = cells.index("gold"), cells.index("silver")
            break
    if gold_col is None:
        raise RuntimeError("LBMA: Gold/Silver columns not found in header")

    months: list[dict] = []
    for row in rows:
        if row is None or len(row) <= max(gold_col, silver_col):
            continue
        g, v = row[gold_col], row[silver_col]
        if not (isinstance(g, (int, float)) and g > 0):
            continue
        if not (isinstance(v, (int, float)) and v > 0):
            continue
        # normalize labels: 'YYYY-MM' strings or datetime cells
        ts = ""
        if row[0] is not None:
            if hasattr(row[0], "strftime"):
                ts = row[0].strftime("%Y-%m")
            else:
                m = re.match(r"(\d{4}-\d{2})", str(row[0]).strip())
                ts = m.group(1) if m else ""
        if not ts:  # label-less newest row → month from the filename
            m = re.search(r"-([A-Za-z]+)-(\d{4})", links[0])
            if m:
                mm = _LBMA_MONTHS.get(m.group(1).lower())
                ts = f"{m.group(2)}-{mm}" if mm else ""
        if not ts:
            continue
        months.append(
            {
                "ts": ts,
                "gold_koz": round(float(g), 1),
                "silver_koz": round(float(v), 1),
                "gold_tonnes": round(koz_to_tonnes(float(g)), 1),
                "silver_tonnes": round(koz_to_tonnes(float(v)), 1),
            }
        )
    if not months:
        raise RuntimeError("LBMA: no data rows")
    latest = months[0]  # file is newest-first; verified by header layout
    return {
        "months": months,
        "latest": latest,
        # legacy keys kept (f2 stored ts/koz/tonnes of silver at top level)
        "ts": latest["ts"],
        "koz": latest["silver_koz"],
        "tonnes": latest["silver_tonnes"],
        "url": links[0],
    }
