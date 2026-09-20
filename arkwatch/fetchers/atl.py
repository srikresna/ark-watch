"""atl.py — Wage Growth Tracker (Atlanta Fed CHCS, XLSX).

URL verified against the official Atlanta Fed page (see the URL constant).
The workbook is multi-sheet; the headline series is the unweighted median
('Overall' column of the data_overall sheet).

FIXED 2026-09-21 (owner health-sweep): the original wiring looked for a
'data_overview' sheet THAT DOES NOT EXIST (the sheet is data_overall) and a
'date' text header that is not there either (row 1 = column labels starting
with an empty cell). The fallback chain landed on data_chart1 and a
first-numeric-after-date scan that grabbed a zero-valued column — 0.0 rows
stored as the wage-growth median, sailing through sanity_min=0.0. The parse
now anchors on the 'Overall' header cell and coerces text numbers (the
sheet ships values as STRINGS, e.g. '4.4').
"""

from __future__ import annotations

import io
from datetime import datetime

import requests
from openpyxl import load_workbook

URL = (
    "https://www.atlantafed.org/-/media/Project/Atlanta/FRBA/Documents/"
    "datafiles/chcs/wage-growth-tracker/wage-growth-data.xlsx"
)
UA = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128.0 Safari/537.36"
}


class AtlError(RuntimeError):
    pass


def _pick_sheet(wb) -> str:
    # Actual sheet names (live-probed 2026-09-21): data_overall, data_chart1..3,
    # WGT_1983, Race, Education, ... — headline = data_overall.
    names = wb.sheetnames
    for cand in ("data_overall", "data_overview"):
        if cand in names:
            return cand
    data_sheets = [n for n in names if str(n).lower().startswith("data")]
    if data_sheets:
        return data_sheets[0]
    raise AtlError(f"atl WGT: data-* sheet not found (available: {names[:8]})")


def _f(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.strip():
        try:
            return float(v.strip())
        except ValueError:
            return None
    return None


def fetch_latest(series_id: str = "ATL:WGT") -> dict:
    r = requests.get(URL, headers=UA, timeout=(10, 60))
    if r.status_code != 200 or r.content[:2] != b"PK":
        raise AtlError(f"atl WGT: HTTP {r.status_code} / not XLSX")
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb[_pick_sheet(wb)]
    rows = list(ws.iter_rows(values_only=True))
    # header row = the one carrying the 'Overall' label; date column is the
    # FIRST column (no 'date' text header exists in this sheet)
    header_i = next(
        (i for i, row in enumerate(rows[:6])
         if row and any(isinstance(c, str) and c.strip().lower() == "overall" for c in row)),
        None,
    )
    if header_i is None:
        raise AtlError("atl WGT: 'Overall' header not found")
    overall_col = next(
        j for j, c in enumerate(rows[header_i])
        if isinstance(c, str) and c.strip().lower() == "overall"
    )
    last = None
    for row in rows[header_i + 1 :]:
        if not row or row[0] is None:
            continue
        v = _f(row[overall_col]) if overall_col < len(row) else None
        if v is not None:
            last = (row[0], v)
    if last is None:
        raise AtlError("atl WGT: no value rows")
    d, v = last
    if isinstance(d, datetime):
        d = d.date().isoformat()
    return {"ts": str(d)[:10], "value": v}
